"""Headless end-to-end CLI workflows using actual BLF, DBC, MF4 and Excel."""
import json
import os
from pathlib import Path
import subprocess
import sys

import numpy as np
import pytest
from asammdf import MDF
from openpyxl import load_workbook

from vehiclelog.models import Formula, Project, Segment
from vehiclelog.project import save_project


ROOT = Path(__file__).resolve().parents[1]


def invoke(*arguments, cache, expected=0):
    env = dict(os.environ, VEHICLELOG_CACHE_DIR=str(cache), PYTHONIOENCODING="utf-8")
    result = subprocess.run([sys.executable, "-m", "vehiclelog", *map(str, arguments)],
                            cwd=ROOT, env=env, capture_output=True, text=True,
                            encoding="utf-8", timeout=60)
    assert result.returncode == expected, result.stdout + "\n" + result.stderr
    return result


@pytest.fixture(scope="module")
def cli_demo(tmp_path_factory):
    directory = tmp_path_factory.mktemp("cli-demo")
    cache = directory / "cache"
    files = json.loads(invoke("demo", "--directory", directory / "samples", cache=cache).stdout)
    inspected = json.loads(invoke("inspect", files["blf"], "--dbc", f"1={files['dbc']}",
                                 "--preview", "2", cache=cache).stdout)
    return files, inspected, cache


def test_cli_demo_inspect_and_cache(cli_demo):
    files, inspected, cache = cli_demo
    assert inspected["metadata"]["frames"] == 2402
    assert inspected["metadata"]["duration"] == 120
    assert len(inspected["metadata"]["frame_preview"]) == 2
    assert {s["name"] for s in inspected["signals"]} == {"PackVoltage", "PackCurrent", "VehicleSpeed"}
    reopened = json.loads(invoke("inspect", files["blf"], "--dbc", f"1={files['dbc']}", cache=cache).stdout)
    assert reopened["cache"] == inspected["cache"]
    assert reopened["metadata"]["frame_preview"] == []


def test_cli_converts_selected_segment_to_mf4_and_excel(cli_demo, tmp_path):
    files, inspected, cache = cli_demo
    mf4 = tmp_path / "regeneration.mf4"
    invoke("convert", files["blf"], mf4, "--format", "mf4", "--dbc", f"1={files['dbc']}",
           "--signal", "PackCurrent", "--segment", "Regeneration:40:80", "--zero-time", cache=cache)
    with MDF(mf4) as mdf:
        signal = mdf.get("PackCurrent")
        assert len(signal.samples) == 401
        assert signal.timestamps[0] == pytest.approx(0)
        assert signal.timestamps[-1] == pytest.approx(40)
        assert signal.samples[0] == -20
        assert signal.samples[-1] == 30  # Exact inclusive sample at stage boundary.
    excel = tmp_path / "regeneration.xlsx"
    invoke("convert", files["blf"], excel, "--format", "excel", "--dbc", f"1={files['dbc']}",
           "--signal", "PackCurrent", "--segment", "Regeneration:40:80", cache=cache)
    workbook = load_workbook(excel, read_only=True, data_only=True)
    try:
        assert {"Summary", "Signals", "Segments", "Formulas", "Results"} <= set(workbook.sheetnames)
        rows = list(workbook["Data_PackCurrent"].values)
        assert len(rows) == 402
        assert rows[1][1:4] == ("Regeneration", 40, -20)
        assert rows[-1][1:4] == ("Regeneration", 80, 30)
    finally:
        workbook.close()


def test_cli_analyzes_only_assigned_stages_and_protects_results(cli_demo, tmp_path):
    files, inspected, cache = cli_demo
    keys = {item["name"]: item["key"] for item in inspected["signals"]}
    stages = [Segment("Discharge", 0, 40), Segment("Regeneration", 40, 80)]
    bindings = {"U": keys["PackVoltage"], "I": keys["PackCurrent"]}
    formulas = [Formula("Discharge energy", "integral(U*I)/3600000", bindings, "kWh", [stages[0].id]),
                Formula("Regenerative energy", "integral(U*I)/3600000", bindings, "kWh", [stages[1].id])]
    project = Project(source=files["blf"], dbc_map={"1": files["dbc"]},
                      segments=stages, formulas=formulas, cache_path=inspected["cache"])
    project_file = save_project(project, tmp_path / "analysis.vls.json")
    destination = tmp_path / "results.json"
    invoke("analyze", project_file, "--output", destination, cache=cache)
    results = json.loads(destination.read_text())
    assert len(results) == 2
    assert [(r["formula_name"], r["segment_name"]) for r in results] == [
        ("Discharge energy", "Discharge"), ("Regenerative energy", "Regeneration")]
    np.testing.assert_allclose([r["value"] for r in results], [400 * 50 * 40 / 3_600_000,
                                                              400 * -20 * 40 / 3_600_000])
    assert {r["status"] for r in results} == {"ok"}
    previous = destination.read_bytes()
    invoke("analyze", project_file, "--output", destination, cache=cache, expected=1)
    assert destination.read_bytes() == previous


def test_cli_reports_unassigned_project(cli_demo, tmp_path):
    files, inspected, cache = cli_demo
    project = Project(source=files["blf"], dbc_map={"1": files["dbc"]},
                      segments=[Segment("Stage", 0, 40)], cache_path=inspected["cache"])
    filename = save_project(project, tmp_path / "unassigned.vls.json")
    result = invoke("analyze", filename, cache=cache, expected=1)
    assert "chưa có công thức" in result.stderr


def test_cli_excel_help_only_offers_native_sampling(tmp_path):
    help_text = invoke("convert", "--help", cache=tmp_path).stdout
    assert "--raster" not in help_text
    assert "--interpolation" not in help_text
    assert "--max-gap" not in help_text
