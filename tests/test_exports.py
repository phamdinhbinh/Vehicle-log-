from contextlib import closing

import can
import numpy as np
import openpyxl
import pytest
from asammdf import MDF
from asammdf.blocks import v4_constants as v4c

from vehiclelog import exports
from vehiclelog.writers import excel, excel_sheets, mf4, raw_mf4
from vehiclelog.demo import create_demo
from vehiclelog.models import CancellationToken, CancelledError, Formula, FormulaResult, Segment, SignalInfo
from vehiclelog.store import StoreWriter


@pytest.fixture
def store(tmp_path):
    source = tmp_path / "source.blf"
    source.write_bytes(b"source placeholder")
    with StoreWriter(tmp_path / "cache", {"source": str(source), "start_time": 1735689600.0,
                                           "warnings": ["example warning"]}) as writer:
        writer.add_signal(SignalInfo("u", "Voltage", "V", channel="1"))
        writer.add_signal(SignalInfo("i", "Current", "A", channel="1"))
        writer.append("u", [0, 1, 2, 3, 10], [400, 401, np.nan, 403, 410])
        writer.append("i", [0, 1, 2, 3, 10], [1, 2, 3, 4, 5])
        return writer.finish()


def _data_rows(book, prefix):
    return [row for name in book.sheetnames if name.startswith(prefix)
            for row in list(book[name].values)[1:]]


def test_excel_native_metadata_formulas_results_and_split(store, tmp_path, monkeypatch):
    monkeypatch.setattr(excel_sheets, "EXCEL_MAX_ROWS", 4)
    monkeypatch.setattr(excel, "CHUNK_SIZE", 2)
    segment = Segment("City", 0, 10)
    formula = Formula("Energy", "integral(U)", {"U": "u"}, "Wh", [segment.id])
    result = FormulaResult(formula.id, formula.name, segment.id, segment.name, None,
                           formula.unit, status="incomplete", message="missing samples", coverage=.5)
    path = exports.export_excel(store, tmp_path / "native.xlsx", segments=[segment],
                                formulas=[formula], results=[result])
    with closing(openpyxl.load_workbook(path, read_only=True, data_only=True)) as book:
        rows = _data_rows(book, "Data_Voltage")
        assert [r[2] for r in rows] == [0, 1, 2, 3, 10]
        assert rows[2][3] is None and rows[2][4] is False
        assert rows[-1][3] == 410
        assert len([s for s in book.sheetnames if s.startswith("Data_Voltage")]) == 2
        assert any(row[0] == "source" for row in _data_rows(book, "Summary"))
        assert _data_rows(book, "Formulas")[0][2] == "integral(U)"
        assert _data_rows(book, "Results")[0][6] == "incomplete"


def test_excel_keeps_distinct_fractional_sample_times_and_gaps(tmp_path, monkeypatch):
    monkeypatch.setattr(excel, "CHUNK_SIZE", 2)
    native = {
        "u": ("Voltage", "V", [.017, .027, .039, .054, 1.123456789], [400, 401, np.nan, 403, 410]),
        "i": ("Current", "A", [.023, .129, .234], [1, 2, 3]),
    }
    with StoreWriter(tmp_path / "unequal-cache", {"duration": 2}) as writer:
        for key, (name, unit, times, values) in native.items():
            writer.add_signal(SignalInfo(key, name, unit))
            writer.append(key, times, values)
        source = writer.finish()
    path = exports.export_excel(source, tmp_path / "native.xlsx")
    with closing(openpyxl.load_workbook(path, read_only=True, data_only=True)) as book:
        assert not any(name.startswith("Raster_") for name in book.sheetnames)
        for name, unit, times, values in native.values():
            rows = _data_rows(book, f"Data_{name}")
            np.testing.assert_allclose([r[2] for r in rows], times, rtol=0, atol=1e-15)
            np.testing.assert_allclose([np.nan if r[3] is None else r[3] for r in rows],
                                       values, equal_nan=True)
            assert len(rows) == len(times)
        assert _data_rows(book, "Data_Voltage")[2][4] is False
        summary = dict(_data_rows(book, "Summary"))
        assert summary["layout"] == "native per signal"
        assert not {"raster_s", "interpolation", "max_gap_s"} & summary.keys()


def test_excel_stage_filters_samples_without_rebasing_or_creating_boundaries(store, tmp_path):
    # Boundaries between recorded samples must not create synthetic rows or reset t=0.
    path = exports.export_excel(store, tmp_path / "stage.xlsx", keys=["i"],
                                segments=[Segment("Cut", .5, 2.5), Segment("Gap", 4, 9)])
    with closing(openpyxl.load_workbook(path, read_only=True, data_only=True)) as book:
        rows = _data_rows(book, "Data_Current")
        assert [r[2] for r in rows] == [1, 2]
        assert [r[3] for r in rows] == [2, 3]
        assert all(r[1] == "Cut" for r in rows)


def test_excel_strings_are_never_formulas(store, tmp_path):
    store.signals["u"].unit = "=1+2"
    path = exports.export_excel(store, tmp_path / "strings.xlsx", keys=["u"])
    with closing(openpyxl.load_workbook(path, read_only=True, data_only=False)) as book:
        assert book["Data_Voltage"]["F2"].value == "=1+2"
        assert book["Data_Voltage"]["F2"].data_type == "s"


def test_decoded_mf4_native_values_invalidity_and_separate_segments(store, tmp_path, monkeypatch):
    monkeypatch.setattr(mf4, "CHUNK_SIZE", 2)
    scopes = [Segment("First", 0, 3), Segment("Second", 9, 10)]
    path = exports.export_mf4(store, tmp_path / "decoded.mf4", keys=["u"], segments=scopes)
    with MDF(path) as mdf:
        assert len(mdf.groups) == 2
        first = mdf.get("Voltage", group=0, ignore_invalidation_bits=True)
        np.testing.assert_array_equal(first.timestamps, [0, 1, 2, 3])
        np.testing.assert_allclose(first.samples, [400, 401, np.nan, 403], equal_nan=True)
        np.testing.assert_array_equal(first.invalidation_bits, [False, False, True, False])
        assert first.unit == "V"
        assert "First" in mdf.groups[0].channel_group.comment
        assert "Second" in mdf.groups[1].channel_group.comment
        assert mdf.header.start_time.timestamp() == 1735689600.0
        np.testing.assert_array_equal(mdf.get("Voltage", group=1).timestamps, [10])


def test_decoded_mf4_zero_time_relative_to_segment_not_first_sample(store, tmp_path):
    path = exports.export_mf4(store, tmp_path / "zero.MF4", keys=["i"],
                              segments=[Segment("Later", 1.5, 3)], zero_time=True)
    with MDF(path) as mdf:
        np.testing.assert_array_equal(mdf.get("Current").timestamps, [.5, 1.5])
        assert '"output_time_offset_s": 1.5' in mdf.groups[0].channel_group.comment


@pytest.mark.parametrize("function,suffix", [(exports.export_excel, ".xlsx"), (exports.export_mf4, ".mf4")])
def test_exports_refuse_existing_or_cache_destinations(store, tmp_path, function, suffix):
    path = tmp_path / f"existing{suffix}"
    path.write_text("keep")
    with pytest.raises(FileExistsError):
        function(store, path)
    assert path.read_text() == "keep"
    with pytest.raises(ValueError, match="bộ nhớ đệm"):
        function(store, store.path / f"bad{suffix}")


@pytest.mark.parametrize("function,suffix", [(exports.export_excel, ".xlsx"), (exports.export_mf4, ".mf4")])
def test_cancelled_exports_remove_temporary_files(store, tmp_path, function, suffix):
    token = CancellationToken()
    destination = tmp_path / f"cancelled{suffix}"
    def progress(fraction, message):
        token.cancel()
    with pytest.raises(CancelledError):
        function(store, destination, progress=progress, token=token)
    assert not destination.exists()
    assert not list(tmp_path.glob(".cancelled-*"))


def test_raw_bus_logging_dbc_extraction_and_streaming(tmp_path, monkeypatch):
    paths = create_demo(tmp_path / "demo")
    monkeypatch.setattr(raw_mf4, "CHUNK_SIZE", 3)
    destination = exports.export_raw_mf4(paths["blf"], tmp_path / "raw.mf4", {1: str(paths["dbc"])},
                                         segments=[Segment("Start", 0, .3)])
    with MDF(destination) as mdf:
        assert mdf.groups[0].channel_group.flags & v4c.FLAG_CG_BUS_EVENT
        assert mdf.groups[0].channel_group.acq_source.bus_type == v4c.BUS_TYPE_CAN
        np.testing.assert_array_equal(mdf.get("CAN_DataFrame.BusChannel").samples, np.ones(8))
        assert set(mdf.get("CAN_DataFrame.ID").samples) == {256, 257}
        assert mdf.get("CAN_DataFrame.DataBytes").samples.shape == (8, 64)
        assert "Start" in mdf.groups[0].channel_group.comment
        with mdf.extract_bus_logging({"CAN": [(paths["dbc"], 1)]}) as decoded:
            voltage = decoded.get("PackVoltage")
            current = decoded.get("PackCurrent")
            np.testing.assert_array_equal(voltage.samples, [400]*4)
            np.testing.assert_array_equal(current.samples, [50]*4)
            np.testing.assert_allclose(voltage.timestamps, [0, .1, .2, .3], atol=2e-7)


def test_raw_retains_extended_fd_flags_dlc_and_channels(tmp_path, monkeypatch):
    source = tmp_path / "flags.blf"
    origin = 1735689600.0
    messages = [can.Message(timestamp=origin, channel=0, arbitration_id=0x123,
                             is_extended_id=False, data=bytes(range(8)), is_rx=True),
                can.Message(timestamp=origin+.25, channel=2, arbitration_id=0x1ABCDE,
                             is_extended_id=True, is_fd=True, bitrate_switch=True,
                             error_state_indicator=True, data=bytes(range(12)), is_rx=False)]
    with can.BLFWriter(source) as writer:
        for message in messages:
            writer.on_message_received(message)
    monkeypatch.setattr(raw_mf4, "CHUNK_SIZE", 1)
    path = exports.export_raw_mf4(source, tmp_path / "flags.mf4")
    with MDF(path) as mdf:
        def values(field):
            return mdf.get("CAN_DataFrame."+field).samples
        np.testing.assert_array_equal(values("BusChannel"), [1, 3])
        np.testing.assert_array_equal(values("ID"), [0x123, 0x1ABCDE])
        for field in ["IDE", "EDL", "BRS", "ESI", "Dir"]:
            np.testing.assert_array_equal(values(field), [0, 1])
        np.testing.assert_array_equal(values("DLC"), [8, 9])
        np.testing.assert_array_equal(values("DataLength"), [8, 12])
        np.testing.assert_array_equal(values("DataBytes")[1, :12], np.arange(12))
        np.testing.assert_allclose(mdf.get("CAN_DataFrame.ID").timestamps, [0, .25])


def test_raw_rejects_remote_frames_and_cleans_output(tmp_path):
    source = tmp_path / "remote.blf"
    with can.BLFWriter(source) as writer:
        writer.on_message_received(can.Message(timestamp=1735689600, arbitration_id=1,
                                                is_remote_frame=True, dlc=8))
    with pytest.raises(ValueError, match="remote/error"):
        exports.export_raw_mf4(source, tmp_path / "bad.mf4")
    assert not (tmp_path / "bad.mf4").exists()
    assert not list(tmp_path.glob(".bad-*"))


def test_mf4_refuses_input_overwrite_even_if_missing_on_disk(store, tmp_path):
    source = tmp_path / "original.mf4"
    store.metadata["source"] = str(source)
    with pytest.raises(ValueError, match="nguồn"):
        exports.export_mf4(store, source)


def test_empty_selection_and_bad_parameters_fail_before_output(store, tmp_path):
    for kwargs in [{"keys": []}, {"keys": ["missing"]}, {"segments": [Segment("Bad", 3, 1)]}]:
        with pytest.raises(ValueError):
            exports.export_excel(store, tmp_path / "bad.xlsx", **kwargs)
    with pytest.raises(ValueError):
        exports.export_mf4(store, tmp_path / "wrong.ext")
    assert not list(tmp_path.glob(".bad-*"))


def test_cancellation_during_mf4_save_removes_temporary_files(store, tmp_path):
    token = CancellationToken()
    def progress(fraction, message):
        if fraction > .92:
            token.cancel()
    with pytest.raises(CancelledError):
        exports.export_mf4(store, tmp_path / "during-save.mf4", token=token, progress=progress)
    assert not (tmp_path / "during-save.mf4").exists()
    assert not list(tmp_path.glob(".during-save-*"))


def test_export_publication_never_overwrites_a_racing_destination(store, tmp_path):
    destination = tmp_path / "race.xlsx"
    def progress(fraction, message):
        if not destination.exists():
            destination.write_text("other task")
    with pytest.raises(FileExistsError):
        exports.export_excel(store, destination, progress=progress)
    assert destination.read_text() == "other task"
    assert not list(tmp_path.glob(".race-*"))


def test_raw_segments_keep_source_origin_and_separate_groups(tmp_path):
    source = tmp_path / "offset.blf"
    # BLF header start timestamps have millisecond resolution; the first frame
    # starts 0.5 ms later. Segment coordinates must use the header origin.
    with can.BLFWriter(source) as writer:
        for timestamp in [1735689600.0005, 1735689600.1005, 1735689600.2005]:
            writer.on_message_received(can.Message(timestamp=timestamp, channel=0,
                                                    arbitration_id=1, is_extended_id=False, data=b"\x01"))
    with can.BLFReader(source) as reader:
        origin = reader.start_timestamp
        relative = [message.timestamp-origin for message in reader]
    segments = [Segment("A", 0, .05), Segment("B", .15, .25)]
    path = exports.export_raw_mf4(source, tmp_path / "cut.mf4", segments=segments, zero_time=True)
    with MDF(path) as mdf:
        assert len(mdf.groups) == 2
        np.testing.assert_allclose(mdf.get("CAN_DataFrame.ID", group=0).timestamps, [relative[0]])
        np.testing.assert_allclose(mdf.get("CAN_DataFrame.ID", group=1).timestamps, [relative[-1]-.15])
        assert '"output_time_offset_s": 0.15' in mdf.groups[1].channel_group.comment


def test_raw_cancellation_cleans_output(tmp_path):
    paths = create_demo(tmp_path / "demo")
    token = CancellationToken()
    with pytest.raises(CancelledError):
        exports.export_raw_mf4(paths["blf"], tmp_path / "cancel-raw.mf4", token=token,
                               progress=lambda fraction, message: token.cancel())
    assert not (tmp_path / "cancel-raw.mf4").exists()
    assert not list(tmp_path.glob(".cancel-raw-*"))


def test_decoded_mf4_includes_only_assigned_pointwise_formulas(store, tmp_path):
    first, second = Segment("A", 0, 1), Segment("B", 1, 3)
    power = Formula("BatteryPower", "U * I / 1000", {"U": "u", "I": "i"}, "kW", [first.id])
    energy = Formula("Energy", "integral(U * I) / 3600000", {"U": "u", "I": "i"}, "kWh", [first.id])
    path = exports.export_mf4(store, tmp_path / "derived.mf4", keys=["i"],
                              segments=[first, second], formulas=[power, energy])
    with MDF(path) as mdf:
        assert len(mdf.groups) == 3
        derived = mdf.get("BatteryPower")
        np.testing.assert_array_equal(derived.timestamps, [0, 1])
        np.testing.assert_allclose(derived.samples, [.4, .802])
        assert derived.unit == "kW"
        assert '"expression": "U * I / 1000"' in derived.comment
        assert "Energy" not in mdf.channels_db
        assert '"name": "A"' in mdf.groups[2].channel_group.comment
    path = exports.export_mf4(store, tmp_path / "implicit-full.mf4", keys=["i"], formulas=[power])
    with MDF(path) as mdf:
        assert "BatteryPower" not in mdf.channels_db
