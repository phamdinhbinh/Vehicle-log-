"""Produce inspectable Excel/MF4 demo outputs and verify their contents."""
from __future__ import annotations

import argparse
from dataclasses import asdict
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import numpy as np
from asammdf import MDF
from openpyxl import load_workbook

from make_demo_project import make_demo_project
from vehiclelog.exports import export_excel, export_mf4, export_raw_mf4
from vehiclelog.formulas import evaluate_all


def verify(directory):
    directory = Path(directory).resolve()
    directory.mkdir(parents=True, exist_ok=False)
    project, store, project_file = make_demo_project(directory / "input")
    results = evaluate_all(store, project.formulas, project.segments)
    assert all(r.status in {"ok", "trace"} for r in results)
    expected = {"Năng lượng ròng": 0.26666666666666666,
                "Năng lượng xả": 0.35555555555555557,
                "Năng lượng thu hồi": 0.08888888888888889,
                "Quãng đường": 1.2, "Mức tiêu thụ": 22.22222222222222}
    for result in results:
        if result.segment_name == "Toàn hành trình" and result.formula_name in expected:
            assert np.isclose(result.value, expected[result.formula_name], rtol=1e-8, atol=1e-10)
    xlsx = export_excel(store, directory / "stage-report.xlsx", segments=project.segments,
                        formulas=project.formulas, results=results, raster=0.1)
    book = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        assert {"Summary", "Signals", "Segments", "Formulas", "Results"}.issubset(book.sheetnames)
        result_rows = list(book["Results"].values)
        assert len(result_rows) == len(results) + 1
        sheetnames = book.sheetnames
    finally:
        book.close()
    decoded = export_mf4(store, directory / "stages-decoded.mf4", segments=project.segments,
                          formulas=project.formulas)
    with MDF(decoded) as mdf:
        voltages = mdf.whereis("PackVoltage")
        assert len(voltages) == 4
        power_groups = mdf.whereis("Công suất pin")
        assert len(power_groups) == 4
        power = mdf.get(group=power_groups[0][0], index=power_groups[0][1])
        assert np.isclose(power.samples[0], 20.0)
        group_count = len(mdf.groups)
    raw = export_raw_mf4(project.source, directory / "can-raw.mf4", dbc_map=project.dbc_map)
    with MDF(raw) as mdf:
        with mdf.extract_bus_logging(database_files={"CAN": [(project.dbc_map["1"], 0)]}) as extracted:
            voltage = extracted.get("PackVoltage")
            np.testing.assert_allclose(voltage.samples, 400)
            assert len(voltage.samples) == 1201
    verification = {"status": "passed", "demo_expected_full_journey": expected,
                    "results": [asdict(r) for r in results],
                    "excel_sheets": sheetnames, "decoded_mf4_groups": group_count,
                    "raw_mf4_dbc_roundtrip": "1201 voltage samples, all 400 V",
                    "files": {"excel": str(xlsx), "decoded": str(decoded), "raw": str(raw), "project": str(project_file)}}
    (directory / "verification.json").write_text(json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8")
    return verification


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--directory", default="examples/verified-output")
    args = parser.parse_args()
    data = verify(args.directory)
    print(json.dumps({k: v for k, v in data.items() if k != "results"}, ensure_ascii=False, indent=2))
